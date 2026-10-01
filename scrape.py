"""Haalt de data van alle teams van D.V. The Pirates (DBMN, via feeds.teambeheer.nl) op en bouwt de site in site/:
een clubpagina plus per team een dashboard en Excel (alleen ruwe data) in site/<slug>/.
Seizoen, team-id en divisie worden automatisch opgezocht.

  python scrape.py                              alle teams met "Pirates" in de naam (clubsite)
  python scrape.py --team "Pirates 7" --solo    alleen dat team, in de root van site/ (voor de pirates7-site)"""
import argparse
import functools
import hashlib
import html as htmllib
import json
import os
import re
import shutil
import traceback
from collections import defaultdict
from pathlib import Path
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from datetime import date, datetime
from zoneinfo import ZoneInfo
from requests.adapters import HTTPAdapter, Retry
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

B = "https://feeds.teambeheer.nl"
D = 41  # DBMN op teambeheer
CLUB_URL = "https://basvanderlit1-commits.github.io/dv-the-pirates/"
PIRATES7_URL = "https://basvanderlit1-commits.github.io/pirates7/"  # eigen site van Pirates 7 (blijft bestaan)
CLUB_TOPIC = "dv-the-pirates-dbmn-uitslagen"  # alle uitslagen van de club
AANVANG, EINDE = "200000", "230000"  # wedstrijden beginnen om 20:00; eindtijd is een schatting (agenda)
# kleuren waarmee teambeheer de plekken in de stand markeert (zelfde regels als de officiële stand)
ZONES = {"label-color-green": "kampioen", "label-color-blue2": "promotie", "raster-column-orange": "nacompetitie",
         "label-color-red": "degradatie"}
HERE = Path(__file__).parent
NOW = datetime.now(ZoneInfo("Europe/Amsterdam"))
http = requests.Session()
http.mount("https://", HTTPAdapter(max_retries=Retry(total=4, backoff_factor=2, status_forcelist=[429, 500, 502, 503, 504])))


@functools.cache  # tegenstanders, standen en klassementen die meerdere teams delen maar één keer ophalen
def _soup(path):
    r = http.get(B + path, timeout=30)
    r.raise_for_status()
    raw = r.content
    try:
        html = raw.decode("utf-8")
    except UnicodeDecodeError:  # sommige namen staan er in cp1252 in
        html = raw.decode("cp1252")
    return BeautifulSoup(html, "html.parser")


def get(path, full=False):
    soup = _soup(path)
    # desktop-weergave; de pagina bevat dezelfde tabellen nogmaals voor tablet/mobiel
    if full:
        return soup
    return next((r for r in soup.select(".computer.only.row") if "mobile" not in r["class"]), soup)


def txt(el):
    return " ".join(el.get_text(" ").split()) if el else ""


def rows(table):
    return [[txt(td) for td in tr.find_all(["td", "th"])] for tr in table.find_all("tr")]


def num(v):
    v = str(v).replace("%", "").strip()
    try:
        return int(v)
    except ValueError:
        try:
            return float(v)
        except ValueError:
            return v


def header_table(soup, first_col):
    for t in soup.find_all("table"):
        h = [txt(th) for th in t.find_all("th")]
        if h and h[0] == first_col:
            return t


def slugify(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")  # "The Pirates 2" -> "the-pirates-2"


def short(name):
    return slugify(name).replace("-", "")  # "Pirates 7" -> "pirates7" (ntfy-topic en Excel-naam, zoals altijd al)


def club_teams():
    """Alle teams met "Pirates" in de naam in het huidige seizoen: naam -> (team-id, seizoen)."""
    teams = {}
    for a in get(f"/web/teams?d={D}", full=True).find_all("a", href=re.compile(r"t=\d+&s=")):
        if "pirates" in txt(a).lower():
            teams.setdefault(txt(a), re.search(r"t=(\d+)&s=([\d-]+)", a["href"]).groups())
    nr = lambda n: int(m.group()) if (m := re.search(r"\d+$", n)) else 99
    return dict(sorted(teams.items(), key=lambda kv: (nr(kv[0]), kv[0])))


def parse_team(page, name):
    """Wedstrijden, spelers en speellocatie van een teampagina (ons team of een tegenstander)."""
    ms = []
    for tr in header_table(page, "#").find("tbody").find_all("tr"):
        td = tr.find_all("td")
        if len(td) < 5:
            continue
        a = td[4].find("a")
        score = txt(a) if a else ""
        home, away = txt(td[2]), txt(td[3])
        res = ""
        if score:
            h, u = map(int, score.split("-"))
            mine, theirs = (h, u) if home == name else (u, h)
            res = "W" if mine > theirs else "V" if mine < theirs else "G"
        opp = (td[3] if home == name else td[2]).find("a")
        ms.append(dict(ronde=txt(td[0]), datum=txt(td[1]), thuis=home, uit=away, score=score,
                       uitslag=res, form=a["href"] if a else None,
                       tegen_id=re.search(r"t=(\d+)", opp["href"]).group(1) if opp else None,
                       vrij=not re.fullmatch(r"\d{2}-\d{2}-\d{4}", txt(td[1]))))  # bv. "Vrije week" in de beker
    players = []
    ptab = header_table(page, "Naam")
    for tr in ptab.find("tbody").find_all("tr") if ptab else []:
        td = tr.find_all("td")
        a = td[0].find("a")
        if not a:
            continue
        rol = txt(td[0].find("b"))
        players.append(dict(naam=txt(a), rol={"C": "Captain", "RC": "Reserve captain"}.get(rol, rol),
                            id=re.search(r"l=(\d+)", a["href"]).group(1), singles=num(txt(td[1])), winst=num(txt(td[2]))))
    loc = page.find(string="Locatie")
    return ms, players, clean_loc(loc.find_parent().find_next("p").stripped_strings) if loc else ""


def clean_loc(parts):
    """Adresregels zonder telefoonnummer, gescheiden door komma's."""
    return ", ".join(p.strip() for p in parts if p.strip() and not re.fullmatch(r"[\d\s\-+()]{6,}", p.strip()))


def game_type(o):
    """Soort partij uit de naam; die verschilt per divisie ("Singel 1 T*", "Koppel 2 - A*", "K 1* - K A", "RR 3 - C* 301")."""
    if o.startswith("RR"):
        return "RR"
    if o.startswith(("Single", "Singel")):
        return "Single"
    if o.startswith("Koppel") or re.match(r"K \d", o):
        return "Koppel"
    return "Team"


def to_date(s):
    try:
        return datetime.strptime(s, "%d-%m-%Y").date()
    except (TypeError, ValueError):
        return s


def sheet(wb, title, header, data):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in data:
        ws.append([num(v) if isinstance(v, str) else v for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, date):
                c.number_format = "dd-mm-yyyy"
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(50, max(len(str(c.value or "")) for c in col) + 2)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def build_team(TEAM, TEAM_ID, S, out, nav):
    """Haalt alles van één team op, schrijft de Excel in out/ en geeft de dashboard-data terug."""
    team = get(f"/web/team?d={D}&t={TEAM_ID}&s={S}")
    DIV = re.search(r"Divisie (\w+)", txt(team.find("a", href=re.compile(r"/web/stand")))).group(1)
    matches, spelers, locatie = parse_team(team, TEAM)

    # ---------- wedstrijdformulieren ----------
    games, bijz = [], []
    for m in matches:
        if not m["form"]:
            continue
        f = get(m["form"])
        thuis = m["thuis"] == TEAM
        for tr in f.find("table").find("tbody").find_all("tr"):
            td = tr.find_all("td")
            sc = txt(td[-1]) if td else ""
            # Totaal ("7 - 2") / niet gespeeld overslaan; "Team 1001/3" (beker) heeft geen spelerskolommen, dus 3 cellen
            if len(td) < 3 or not re.fullmatch(r"\d+-\d+", sc):
                continue
            hp = [txt(a) for a in td[1].find_all("a")]
            up = [txt(a) for a in td[2].find_all("a")]
            h, u = map(int, sc.split("-"))
            mine, theirs = (h, u) if thuis else (u, h)
            games.append(dict(ronde=m["ronde"], datum=m["datum"], tegen=m["uit"] if thuis else m["thuis"],
                              onderdeel=txt(td[0]), wij=", ".join(hp if thuis else up) or "(team)",
                              zij=", ".join(up if thuis else hp) or "(team)", legs_wij=mine, legs_zij=theirs,
                              uitslag="W" if mine > theirs else "V"))
        loc = f.find(string="Locatie")
        m["locatie"] = clean_loc(" | ".join(loc.find_parent().find_next("a").parent.stripped_strings)
                                 .split(" | Bijz.")[0].removeprefix("Locatie | ").split(" | ")) if loc else ""
        for item in f.select(".ui.list .item"):
            bijz.append(dict(ronde=m["ronde"], datum=m["datum"], speler=txt(item.select_one(".header")),
                             prestatie=txt(item.select_one(".content")).replace(txt(item.select_one(".header")), "", 1).strip()))
    for m in matches:  # controle: uitslag = som van de partijen, waarbij de round robin samen één punt is
        if m["uitslag"]:
            gs = [g for g in games if g["ronde"] == m["ronde"]]
            rr = [g for g in gs if game_type(g["onderdeel"]) == "RR"]
            rw, rv = sum(g["legs_wij"] for g in rr), sum(g["legs_zij"] for g in rr)
            som = (sum(g["uitslag"] == "W" for g in gs if g not in rr) + (rw > rv), sum(g["uitslag"] == "V" for g in gs if g not in rr) + (rv > rw))
            h, u = map(int, m["score"].split("-"))
            if som != ((h, u) if m["thuis"] == TEAM else (u, h)):
                print(f"LET OP ({TEAM}) ronde {m['ronde']}: uitslag {m['score']} (thuis-uit), partijen tellen op tot {som[0]}-{som[1]} (wij-zij)")
    if onbekend :={g["onderdeel"] for g in games if game_type(g["onderdeel"]) == "Team" and not g["onderdeel"].startswith("Team")}:
        print(f"LET OP ({TEAM}): onderdelen niet herkend als single/koppel/RR, tellen niet mee per speler:", onbekend)
    ours = {s["naam"] for s in spelers}
    for b in bijz:
        b["team"] = TEAM if b["speler"] in ours else "tegenstander"

    # ---------- per speler: statistiek uit wedstrijdformulieren + spelerpagina ----------
    agg = defaultdict(lambda: defaultdict(int))
    for g in games:
        kind = {"Single": "s", "Koppel": "k", "RR": "rr"}.get(game_type(g["onderdeel"]))
        if not kind:
            continue
        for p in g["wij"].split(", "):
            a = agg[p]
            a[kind + "_gesp"] += 1
            a[kind + "_w"] += g["uitslag"] == "W"
            a[kind + "_lv"] += g["legs_wij"]
            a[kind + "_lt"] += g["legs_zij"]
    for s in spelers:
        sp = get(f"/web/speler?d={D}&l={s['id']}&s={S}")
        stats = {txt(st.select_one(".label")): txt(st.select_one(".value")) for st in sp.select(".statistic")}
        s["site_singles_pct"] = next((v for k, v in stats.items() if "singles" in k), "")
        s["site_koppels_pct"] = next((v for k, v in stats.items() if "koppels" in k), "")

    # ---------- stand, klassementen, bijzondere resultaten van de poule ----------
    stand_soup = get(f"/web/stand?d={D}&s={S}&div={DIV}", full=True)
    stand_tab = next(t for t in stand_soup.find_all("table") if "Wed" in [txt(th) for th in t.find_all("th")])
    stand = rows(stand_tab)
    zones = {txt(tds[1]): next((ZONES[c] for c in tds[0].get("class") or [] if c in ZONES), "")
             for tds in (tr.find_all("td") for tr in stand_tab.find_all("tr")) if len(tds) > 1}

    pk_single = rows(get(f"/web/scorelijst-pk/?d={D}&mt=1&s={S}&filter=P-{DIV}").find("table"))
    pk_koppel = rows(get(f"/web/scorelijst-pk/?d={D}&mt=2&s={S}&filter=P-{DIV}").find("table"))
    # punten per ronde -> onze positie na elke gespeelde ronde (benadering: telt alle gespeelde punten tot en met die ronde)
    rs = next(rows(t) for t in stand_soup.find_all("table") if "Wed" not in [txt(th) for th in t.find_all("th")])
    rcols = [i for i, h in enumerate(rs[0]) if h.isdigit()]
    pts = {r[1]: [num(r[i]) if i < len(r) else "" for i in rcols] for r in rs[1:] if len(r) > 2}
    cum, verloop = defaultdict(float), []
    for k, i in enumerate(rcols):
        for t, p in pts.items():
            cum[t] += p[k] if isinstance(p[k], (int, float)) else 0
        if isinstance(pts.get(TEAM, [None] * len(rcols))[k], (int, float)):
            verloop.append(dict(ronde=int(rs[0][i]), pos=1 + sum(v > cum[TEAM] for t, v in cum.items() if t != TEAM)))

    # tegenstanders: adres (route), vorm en beste spelers
    tegenstanders = {}
    for m in matches:
        naam = m["uit"] if m["thuis"] == TEAM else m["thuis"]
        if m["tegen_id"] and naam not in tegenstanders:
            o_ms, o_sp, o_loc = parse_team(get(f"/web/team?d={D}&t={m['tegen_id']}&s={S}"), naam)
            top = [p for p in o_sp if isinstance(p["singles"], int) and p["singles"] > 0 and isinstance(p["winst"], (int, float)) and p["winst"] > 0]
            tegenstanders[naam] = dict(locatie=o_loc, vorm=[x["uitslag"] for x in o_ms if x["uitslag"]][-5:],
                                       spelers=[dict(naam=p["naam"], singles=p["singles"], winst=p["winst"])
                                                for p in sorted(top, key=lambda p: (-p["winst"], -p["singles"]))[:3]])
    for m in matches:  # locatie van nog te spelen wedstrijden: thuis bij ons, uit bij de tegenstander
        if not m.get("locatie") and not m["vrij"]:
            m["locatie"] = locatie if m["thuis"] == TEAM else tegenstanders.get(m["uit"] if m["thuis"] == TEAM else m["thuis"], {}).get("locatie", "")

    bijz_lists = {}
    for t, name in [(1, "180ers"), (2, "Hoogste finishes"), (3, "Snelste leg"), (4, "171ers")]:
        tab = get(f"/web/scorelijst-bijzres/?d={D}&t={t}&s={S}&filter=P-{DIV}").find("table")
        bijz_lists[name] = rows(tab) if tab else [["(geen data)"]]

    # ---------- voorbereiden ----------
    for m in matches:
        thuis = m["thuis"] == TEAM
        m["tegen"], m["tu"] = (m["uit"], "Thuis") if thuis else (m["thuis"], "Uit")
        if m["score"]:
            h, u = m["score"].split("-")
            m["wijzij"] = f"{h} - {u}" if thuis else f"{u} - {h}"
    role = lambda r: ", ".join(s["naam"] for s in spelers if s["rol"] == r) or "–"
    ours_only = lambda tab, col: [r for r in tab[1:] if len(r) > col and r[col] == TEAM]
    pk_count = lambda tab: sum(1 for r in tab[1:] if len(r) > 2)  # aantal spelers in het klassement
    won = lambda r: round(num(r[3]) * num(r[6]) / 100)  # gewonnen partijen = gespeeld x winst%
    kind = game_type
    out.mkdir(parents=True, exist_ok=True)
    xlsx = f"{short(TEAM)}-resultaten.xlsx"

    # ---------- Excel: alleen de uitgelezen data (de presentatie zit in het dashboard) ----------
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "Info", ["Veld", "Waarde"],
          [["Team", TEAM], ["Seizoen", S], ["Divisie", DIV], ["Speellocatie", locatie],
           ["Captain", role("Captain")], ["Reserve captain", role("Reserve captain")],
           ["Bijgewerkt", NOW.strftime("%d-%m-%Y %H:%M")], ["Bron", f"{B}/web/team?d={D}&t={TEAM_ID}&s={S}"]])
    sheet(wb, "Stand", ["Positie", "Team", "Gespeeld", "Winst", "Verlies", "Punten", "Gemiddeld", "Strafpunten"],
          [r for r in stand[1:] if any(r)])
    sheet(wb, "Programma", ["Ronde", "Datum", "Thuis", "Uit", "Tegenstander", "Thuis/Uit", "Punten wij", "Punten zij",
                            "Uitslag", "Vrije week", "Locatie", "Wedstrijdformulier"],
          [[m["ronde"], to_date(m["datum"]), m["thuis"], m["uit"], m["tegen"], m["tu"],
            int(m["wijzij"].split(" - ")[0]) if m["score"] else None, int(m["wijzij"].split(" - ")[1]) if m["score"] else None,
            m["uitslag"] or None, "ja" if m["vrij"] else None, m.get("locatie", "") or None,
            B + m["form"] if m["form"] else None] for m in matches])
    sheet(wb, "Partijen", ["Ronde", "Datum", "Tegenstander", "Onderdeel", "Type", TEAM, "Tegenstander(s)",
                           "Legs wij", "Legs zij", "Uitslag"],
          [[g["ronde"], to_date(g["datum"]), g["tegen"], g["onderdeel"], kind(g["onderdeel"]), g["wij"], g["zij"],
            g["legs_wij"], g["legs_zij"], g["uitslag"]] for g in games])
    sheet(wb, "Spelers", ["Speler", "Rol",
                          "Singles gespeeld", "Singles gewonnen", "Singles legs voor", "Singles legs tegen",
                          "Koppels gespeeld", "Koppels gewonnen", "Koppels legs voor", "Koppels legs tegen",
                          "RR gespeeld", "RR gewonnen", "RR legs voor", "RR legs tegen",
                          "180's", "Finishes", "Winst % singles (site)", "Winst % koppels (site)"],
          [[s["naam"], s["rol"] or "Speler",
            *(agg[s["naam"]][k + x] for k in ("s", "k", "rr") for x in ("_gesp", "_w", "_lv", "_lt")),
            sum("180" in b["prestatie"] for b in bijz if b["speler"] == s["naam"]),
            ", ".join(b["prestatie"].replace(" finish", "") for b in bijz
                      if b["speler"] == s["naam"] and "finish" in b["prestatie"]) or None,
            s["site_singles_pct"] or None, s["site_koppels_pct"] or None] for s in spelers])
    for label, tab in [("Singles", pk_single), ("Koppels", pk_koppel)]:
        sheet(wb, f"Klassement {label.lower()}", ["Positie", "Spelers in klassement", "Speler", "Gespeeld", "Gewonnen", "Winst %"],
              [[r[0], pk_count(tab), r[1], r[3], won(r), r[6]] for r in ours_only(tab, 2)])
    sheet(wb, "Bijzondere resultaten", ["Ronde", "Datum", "Speler", "Prestatie"],
          [[b["ronde"], to_date(b["datum"]), b["speler"], b["prestatie"]] for b in bijz if b["team"] == TEAM])
    sheet(wb, "Klassering bijzonder", ["Lijst", "Positie", "Speler", "Aantal / waarde"],
          [[name, r[0], r[1], r[4]] for name, tab in bijz_lists.items() for r in ours_only(tab, 2)])
    wb.save(out / xlsx)

    # ---------- data voor het dashboard ----------
    print("OK", TEAM, S, DIV, "-", len(matches), "wedstrijden,", len(games), "partijen,", len(spelers), "spelers")
    return dict(
        team=TEAM, slug=slugify(TEAM), div=DIV, seizoen=S, bijgewerkt=NOW.strftime("%d-%m-%Y %H:%M"), locatie=locatie,
        captain=role("Captain"), rc=role("Reserve captain"), bron=f"{B}/web/team?d={D}&t={TEAM_ID}&s={S}",
        stand=[dict(pos=num(r[0]), team=r[1], wed=num(r[2]), w=num(r[3]), v=num(r[4]), pnt=num(r[5]), gem=num(r[6]), zone=zones.get(r[1], ""))
               for r in stand[1:] if any(r)],
        matches=[dict(ronde=m["ronde"], datum=m["datum"], tegen=m["tegen"], tu=m["tu"], uitslag=m["uitslag"], vrij=m["vrij"],
                      wij=int(m["wijzij"].split(" - ")[0]) if m["score"] else None,
                      zij=int(m["wijzij"].split(" - ")[1]) if m["score"] else None,
                      locatie=m.get("locatie", ""), form=B + m["form"] if m["form"] else None)
                 for m in matches],
        games=[dict(ronde=g["ronde"], type=kind(g["onderdeel"]), onderdeel=g["onderdeel"], wij=g["wij"], zij=g["zij"],
                    lw=g["legs_wij"], lz=g["legs_zij"], uitslag=g["uitslag"]) for g in games],
        spelers=[dict(naam=s["naam"], rol=s["rol"] or "Speler",
                      **{k: [agg[s["naam"]][k + "_gesp"], agg[s["naam"]][k + "_w"], agg[s["naam"]][k + "_lv"],
                             agg[s["naam"]][k + "_lt"]] for k in ("s", "k", "rr")},
                      n180=sum("180" in b["prestatie"] for b in bijz if b["speler"] == s["naam"]),
                      finishes=[b["prestatie"].replace(" finish", "") for b in bijz
                                if b["speler"] == s["naam"] and "finish" in b["prestatie"]]) for s in spelers],
        pk={label: [dict(pos=num(r[0]), naam=r[1], n=num(r[3]), w=won(r), pct=num(r[6]))
                    for r in ours_only(tab, 2)] for label, tab in [("Singles", pk_single), ("Koppels", pk_koppel)]},
        pk_totaal={"Singles": pk_count(pk_single), "Koppels": pk_count(pk_koppel)},
        bijz=[dict(ronde=b["ronde"], datum=b["datum"], speler=b["speler"], prestatie=b["prestatie"])
              for b in bijz if b["team"] == TEAM],
        bijzrank=[dict(lijst=name, pos=num(r[0]), speler=r[1], waarde=num(r[4]))
                  for name, tab in bijz_lists.items() for r in ours_only(tab, 2)],
        verloop=verloop, tegenstanders=tegenstanders, ntfy=f"{short(TEAM)}-dbmn-uitslagen", xlsx=xlsx, ics=f"{short(TEAM)}.ics",
        club=CLUB_URL, clubtopic=CLUB_TOPIC, nav=nav,
    )


def write_page(out, template, data, name, short_name, desc):
    """index.html uit een template, plus de bestanden uit web/ met een manifest dat bij deze map hoort."""
    out.mkdir(parents=True, exist_ok=True)
    html = (HERE / template).read_text(encoding="utf-8") \
        .replace("/*DATA*/null", json.dumps(data, ensure_ascii=False).replace("</", "<\\/")) \
        .replace("__NAME__", htmllib.escape(name)).replace("__SHORT__", htmllib.escape(short_name))
    web = list((HERE / "web").iterdir())
    for f in web:
        shutil.copy(f, out / f.name)
    man = json.loads((HERE / "web" / "manifest.webmanifest").read_text(encoding="utf-8"))
    man.update(name=name, short_name=short_name, description=desc)  # start_url en scope "./" = deze map
    (out / "manifest.webmanifest").write_text(json.dumps(man, ensure_ascii=False, indent=2), encoding="utf-8")
    for f in web:  # versienummer achter elk bestand, zodat browsers na een wijziging nooit een oude kopie tonen
        html = html.replace(f'"{f.name}"', f'"{f.name}?v={hashlib.md5((out / f.name).read_bytes()).hexdigest()[:8]}"')
    (out / "index.html").write_text(html, encoding="utf-8")


def notify(d, oud, click):
    """Melding bij elke nieuwe uitslag, naar het topic van het team en naar het clubtopic (ntfy.sh)."""
    me = next((s for s in d["stand"] if s["team"] == d["team"]), {})
    for m in d["matches"]:
        if m["uitslag"] and m["ronde"] not in oud:
            score = f"{m['wij']}-{m['zij']}"
            titel = {"W": f"{d['team']} wint {score} van {m['tegen']}", "V": f"{d['team']} verliest {score} van {m['tegen']}",
                     "G": f"{d['team']} speelt {score} gelijk tegen {m['tegen']}"}[m["uitslag"]]
            wat = "Beker" if m["ronde"].startswith("b") else "Ronde " + m["ronde"]
            tekst = f"{wat} · {m['datum']} · {m['tu'].lower()}\nStand: {me.get('pos')}e in {d['div']} met {me.get('pnt')} punten"
            for topic in (d["ntfy"], CLUB_TOPIC):
                try:  # een storing bij ntfy mag de update van het dashboard nooit tegenhouden
                    http.post("https://ntfy.sh/", json=dict(topic=topic, title=titel, message=tekst, click=click,
                                                          tags=["dart", "trophy"] if m["uitslag"] == "W" else ["dart"]),
                              timeout=30).raise_for_status()
                    print("Melding verstuurd:", topic, titel)
                except requests.RequestException as e:
                    print("Melding mislukt:", topic, e)


def keep_previous(out, naam, url):
    """Faalt een team, dan de laatste goede versie van de live site overnemen (pagina, Excel, agenda)."""
    out.mkdir(parents=True, exist_ok=True)
    namen = ["index.html", f"{short(naam)}-resultaten.xlsx", f"{short(naam)}.ics"]
    try:  # eerst alles ophalen, pas schrijven als alles lukt (geen half overgenomen pagina)
        inhoud = {f: http.get(url + f, timeout=30) for f in namen}
        for r in inhoud.values():
            r.raise_for_status()
    except requests.RequestException as e:
        print("Vorige versie niet op te halen:", naam, e)
        for f in namen:  # half geschreven uitvoer van de mislukte build weghalen
            (out / f).unlink(missing_ok=True)
        return False
    for f, r in inhoud.items():
        (out / f).write_bytes(r.content)
    for f in (HERE / "web").iterdir():  # afbeeldingen, stijl en scripts waar de pagina naar verwijst
        if not (out / f.name).exists():
            shutil.copy(f, out / f.name)
    return True


def alarm(titel, tekst):
    """Seintje aan de beheerder (privé ntfy-topic uit NTFY_BEHEER); zonder topic alleen in de log."""
    print("ALARM:", titel, "-", tekst)
    if topic := os.environ.get("NTFY_BEHEER"):
        try:
            http.post("https://ntfy.sh/", json=dict(topic=topic, title=titel, message=tekst, priority=4, tags=["warning"],
                                                  click=os.environ.get("RUN_URL", CLUB_URL)), timeout=30).raise_for_status()
        except requests.RequestException as e:
            print("Alarm versturen mislukt:", e)


def summary(d, url):
    """Wat de clubpagina van een team nodig heeft (ook de vorige stand voor meldingen)."""
    me = next((s for s in d["stand"] if s["team"] == d["team"]), {})
    return dict(naam=d["team"], url=url, div=d["div"], pos=me.get("pos"), teams=len(d["stand"]), pnt=me.get("pnt"),
                gem=me.get("gem"), w=me.get("w"), v=me.get("v"), ntfy=d["ntfy"],
                matches=[{k: m[k] for k in ("ronde", "datum", "tegen", "tu", "uitslag", "wij", "zij", "vrij", "locatie")}
                         for m in d["matches"]],
                bijz=d["bijz"], zone=me.get("zone", ""), ics=d["ics"], locatie=d["locatie"], spelers=[s["naam"] for s in d["spelers"]],
                games=[[g["ronde"], g["type"], g["wij"], g["uitslag"], g["lw"], g["lz"]] for g in d["games"]])


def fold(line):
    """iCalendar: regels van hooguit 75 bytes, vervolgregels beginnen met een spatie."""
    b, out = line.encode(), []
    while len(b) > (75 if not out else 74):
        cut = 75 if not out else 74
        while b[cut] & 0xC0 == 0x80:  # niet midden in een UTF-8-teken knippen
            cut -= 1
        out.append(b[:cut])
        b = b[cut:]
    return "\r\n ".join(x.decode() for x in out + [b])


def ics(naam, teams):
    """Agenda (iCalendar) met alle wedstrijden; aanvang 20:00 (vast bij de DBMN, niet op teambeheer), einde geschat."""
    esc = lambda s: str(s).replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")
    stamp = NOW.astimezone(ZoneInfo("UTC")).strftime("%Y%m%dT%H%M%SZ")
    L = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//D.V. The Pirates//dv-the-pirates//NL", "CALSCALE:GREGORIAN",
         "METHOD:PUBLISH", f"X-WR-CALNAME:{esc(naam)}", "X-WR-TIMEZONE:Europe/Amsterdam",
         "REFRESH-INTERVAL;VALUE=DURATION:PT6H", "X-PUBLISHED-TTL:PT6H",
         # tijdzone meesturen, zodat 20:00 in zomer- en wintertijd klopt
         "BEGIN:VTIMEZONE", "TZID:Europe/Amsterdam",
         "BEGIN:DAYLIGHT", "TZOFFSETFROM:+0100", "TZOFFSETTO:+0200", "TZNAME:CEST", "DTSTART:19700329T020000",
         "RRULE:FREQ=YEARLY;BYMONTH=3;BYDAY=-1SU", "END:DAYLIGHT",
         "BEGIN:STANDARD", "TZOFFSETFROM:+0200", "TZOFFSETTO:+0100", "TZNAME:CET", "DTSTART:19701025T030000",
         "RRULE:FREQ=YEARLY;BYMONTH=10;BYDAY=-1SU", "END:STANDARD", "END:VTIMEZONE"]
    for d, url in teams:
        for m in d["matches"]:
            if m["vrij"]:
                continue
            dag = datetime.strptime(m["datum"], "%d-%m-%Y").date()
            wat = "Beker" if m["ronde"].startswith("b") else f"Ronde {m['ronde']}"
            titel = f"{d['team']} {m['wij']}-{m['zij']} {m['tegen']}" if m["uitslag"] else f"{d['team']} - {m['tegen']}"
            L += ["BEGIN:VEVENT", f"UID:{slugify(d['team'])}-{m['ronde']}@dv-the-pirates", f"DTSTAMP:{stamp}",
                  f"DTSTART;TZID=Europe/Amsterdam:{dag:%Y%m%d}T{AANVANG}", f"DTEND;TZID=Europe/Amsterdam:{dag:%Y%m%d}T{EINDE}",
                  f"SUMMARY:{esc(titel + ' (' + m['tu'].lower() + ')')}", f"LOCATION:{esc(m['locatie'])}",
                  f"DESCRIPTION:{esc(wat + ' · DBMN divisie ' + d['div'] + chr(10) + url)}", f"URL:{url}",
                  "END:VEVENT"]
    return "\r\n".join(fold(x) for x in L + ["END:VCALENDAR"]) + "\r\n"


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--team", action="append", help="alleen dit team (mag vaker); standaard alle Pirates-teams")
    ap.add_argument("--solo", action="store_true", help="één team in de root van site/, zonder clubpagina")
    args = ap.parse_args()
    teams = club_teams()
    chosen = args.team or list(teams)
    if missing := [t for t in chosen if t not in teams]:
        raise SystemExit(f"Team {missing} niet gevonden op {B}/web/teams?d={D} - klopt de naam nog?")
    nav = [dict(naam=n, url=CLUB_URL + slugify(n) + "/") for n in teams]
    site = HERE / "site"
    shutil.rmtree(site, ignore_errors=True)  # teams die niet meer bestaan niet laten staan

    prev = Path("prev.json")  # vorige clubstand (GitHub Actions haalt de live data.json op)
    try:
        vorige = {t["naam"]: t for t in json.loads(prev.read_text(encoding="utf-8"))["teams"]} if prev.exists() else {}
    except (ValueError, KeyError, TypeError) as e:  # kapotte of oude vorige stand: dan alleen geen meldingen
        print("LET OP: prev.json onleesbaar, geen meldingen deze keer:", e)
        vorige = {}
    oud = {n: {m["ronde"] for m in t["matches"] if m["uitslag"]} for n, t in vorige.items()}

    club, teamdata, fouten = [], [], []
    for naam in chosen:
        team_id, s = teams[naam]
        out = site if args.solo else site / slugify(naam)
        url = PIRATES7_URL if args.solo else CLUB_URL + slugify(naam) + "/"
        try:
            d = build_team(naam, team_id, s, out, nav)
            write_page(out, "dashboard_template.html", d, f"{naam} Dashboard", naam,
                       f"Stand, uitslagen en statistieken van {naam} (DBMN Divisie {d['div']})")
            (out / d["ics"]).write_bytes(ics(naam, [(d, url)]).encode())
            samenvatting = summary(d, url)
        except Exception as e:  # één afwijkende teampagina mag de rest van de site niet tegenhouden
            traceback.print_exc()
            bewaard = keep_previous(out, naam, url)
            fouten.append(f"{naam}: {type(e).__name__}: {e}"[:300] + ("" if bewaard else " (geen vorige versie)"))
            if naam in vorige and not args.solo:  # clubpagina en clubagenda: de vorige stand van dit team
                club.append(vorige[naam])
                teamdata.append((dict(team=naam, div=vorige[naam]["div"], matches=vorige[naam]["matches"]), url))
            continue
        if os.environ.get("MELDINGEN") == "aan" and naam in oud:  # nieuw team: eerst een vorige stand opbouwen
            notify(d, oud[naam], url)
        club.append(samenvatting)
        teamdata.append((d, url))
    if fouten:
        alarm(f"{len(fouten)} van {len(chosen)} teams niet bijgewerkt",
              "Deze teams tonen de vorige stand:\n" + "\n".join(fouten))
        if len(fouten) == len(chosen):
            raise SystemExit("Geen enkel team bijgewerkt - teambeheer onbereikbaar of veranderd?")

    if not args.solo:
        data = dict(bijgewerkt=NOW.strftime("%d-%m-%Y %H:%M"), seizoen=teams[chosen[0]][1], teams=club,
                    ntfy=CLUB_TOPIC, pirates7=PIRATES7_URL, bron=f"{B}/web/teams?d={D}", ics="dv-the-pirates.ics",
                    zones={v: v for v in ZONES.values()})
        (site / "dv-the-pirates.ics").write_bytes(ics("D.V. The Pirates", teamdata).encode())
        (site / "data.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        write_page(site, "club_template.html", data, "D.V. The Pirates", "The Pirates",
                   "Alle teams van D.V. The Pirates in de DBMN: stand, uitslagen en programma")
        print("OK clubpagina -", len(club), "teams")
